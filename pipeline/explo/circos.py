"""Spike étape 1 : code postal → circonscription(s) législative(s) → député (XVIIe).

Script jetable. Lit les fichiers téléchargés dans data/raw/circos/ (non fiables :
on ne fait que les parser) et affiche les mesures du rapport. Écrit quelques
tables dérivées dans --out (par défaut : ../data/raw/circos/derive/).

    cd pipeline && uv run python explo/circos.py
    cd pipeline && uv run --with openpyxl --with xlrd python explo/circos.py   # + tables MI 2012/2017

Sources (voir le rapport pour les URLs) :
- mi2024_t1_bureaux.csv / mi2024_t1_circos.csv : résultats définitifs du 1er tour
  2024 (ministère de l'Intérieur). Le fichier par bureau n'a PAS de colonne
  circonscription : on la retrouve en appariant la liste des candidats du bureau
  avec celle de la circonscription (même département).
- mi2022_t1_burvot.txt / mi2022_t1_subcom.txt : résultats du 1er tour 2022, qui
  portent le code de circonscription.
- laposte_hexasmal.csv : base officielle des codes postaux (La Poste).
- cog2026_*.csv : Code officiel géographique 2026 (Insee).
- insee_reu_bv.parquet / insee_reu_adresses.parquet : bureaux de vote et adresses
  de leurs électeurs (Insee, REU 2022-2023), optionnels.
- mi2012_correspondance.xls / mi2017_correspondance.xlsx : tables de correspondance
  communes × cantons × circos du ministère de l'Intérieur, optionnelles.
- dump AMO de l'AN (acteurs) : jointure circonscription → député.
"""

from __future__ import annotations

import argparse
import collections
import glob
import io
import json
from pathlib import Path

import polars as pl

ICI = Path(__file__).resolve().parent
RACINE = ICI.parent.parent
RAW = RACINE / "data" / "raw" / "circos"
AMO = RACINE / "data" / "raw" / "an_amo" / "x" / "json" / "acteur"

# Codes département du ministère de l'Intérieur → numDepartement de l'AN.
DEP_MI = {
    "ZA": "971", "ZB": "972", "ZC": "973", "ZD": "974", "ZM": "976",
    "ZS": "975", "ZX": "977", "ZN": "988", "ZP": "987", "ZW": "986",
    "ZZ": "099",
}
# Préfixe du code commune Insee pour les codes MI à 3 chiffres hors métropole.
PREFIXE_OM = {"ZA": "97", "ZB": "97", "ZC": "97", "ZD": "97", "ZM": "97",
              "ZS": "97", "ZX": "97", "ZN": "98", "ZP": "98", "ZW": "98"}
# Codes commune du MI 2024 sans équivalent direct au COG (vérifiés à la main).
CODES_SPECIAUX = {
    # Partie axonaise (ex-Gernicourt, 02344) de la commune nouvelle de Cormicy (Marne),
    # restée dans la 1re circo de l'Aisne.
    "02899": ["51171"],
    # Wallis-et-Futuna : une seule « commune » au MI, trois circonscriptions
    # territoriales au COG (Alo, Sigave, Uvea).
    "98601": ["98611", "98612", "98613"],
}
# Communes à arrondissements municipaux : code commune → préfixe de l'arrondissement.
PLM = {"75056": ("751", "Paris"), "13055": ("132", "Marseille"),
       "69123": ("6938", "Lyon")}


def titre(t: str) -> None:
    print(f"\n{'=' * 78}\n{t}\n{'=' * 78}")


def lire_csv(path: Path, latin1: bool = False, **kw) -> pl.DataFrame:
    data = path.read_bytes()
    if latin1:
        data = data.decode("latin-1").encode("utf-8")
    return pl.read_csv(io.BytesIO(data), separator=";", infer_schema=False, **kw)


def dep_an(dep_mi: str) -> str:
    dep_mi = dep_mi.strip()
    if dep_mi in DEP_MI:
        return DEP_MI[dep_mi]
    if dep_mi.isdigit() and len(dep_mi) < 2:
        return dep_mi.zfill(2)
    return dep_mi


def circo_id(dep: str, num) -> str:
    return f"{dep}-{int(num):02d}"


def arrondissement(code_commune: str, bv: str) -> str:
    """Paris, Marseille, Lyon : le n° de bureau commence par le n° d'arrondissement."""
    prefixe, _ = PLM[code_commune]
    arr = int(bv) // 100
    if arr == 0:
        return code_commune + "?"  # bureau hors numérotation (ex. Lyon « 0001 »)
    return prefixe + (f"{arr:02d}" if len(prefixe) == 3 else f"{arr}")


# --------------------------------------------------------------------------
# 2024 : circonscription de chaque bureau par appariement des candidats
# --------------------------------------------------------------------------

def cle_candidats(df: pl.DataFrame) -> pl.Expr:
    n = max(int(c.rsplit(" ", 1)[1]) for c in df.columns if c.startswith("Nom candidat "))
    parts = [
        pl.concat_str([pl.col(f"Numéro de panneau {i}"), pl.col(f"Nom candidat {i}"),
                       pl.col(f"Prénom candidat {i}")], separator=":")
        for i in range(1, n + 1)
    ]
    return pl.concat_list(parts).list.drop_nulls().list.sort().list.join("|")


def bureaux_2024(raw: Path) -> pl.DataFrame:
    circos = lire_csv(raw / "mi2024_t1_circos.csv")
    circos = circos.with_columns(
        dep=pl.col("Code département").map_elements(dep_an, return_dtype=pl.String),
        num=pl.col("Code circonscription législative").str.slice(-2).cast(pl.Int32),
        cle=cle_candidats(circos),
    )
    circos = circos.with_columns(
        circo=pl.concat_str([pl.col("dep"), pl.col("num").cast(pl.String).str.zfill(2)],
                            separator="-")
    ).select("dep", "circo", "cle")
    assert circos.height == 577, circos.height
    assert circos["cle"].n_unique() == 577, "clé candidats non unique"

    bv = lire_csv(raw / "mi2024_t1_bureaux.csv")
    bv = bv.with_columns(
        dep=pl.col("Code département").map_elements(dep_an, return_dtype=pl.String),
        cle=cle_candidats(bv),
        inscrits=pl.col("Inscrits").cast(pl.Int64),
    )
    bv = bv.join(circos, on=["dep", "cle"], how="left")
    sans = bv.filter(pl.col("circo").is_null())
    print(f"2024 : {bv.height} bureaux, {sans.height} sans circonscription retrouvée")
    if sans.height:
        print(sans.select("Code département", "Code commune", "Libellé commune", "Code BV").head(10))
    return bv.select(
        pl.col("dep"), pl.col("Code commune").alias("com_mi"),
        pl.col("Libellé commune").alias("libelle"), pl.col("Code BV").alias("bv"),
        "inscrits", "circo",
    )


def code_insee_2024(com_mi: str) -> str | None:
    if com_mi.startswith("ZZ"):
        return None
    if com_mi.startswith("ZX"):
        return "97" + com_mi[2:]
    return com_mi.zfill(5)


# --------------------------------------------------------------------------
# 2022 : la circonscription est dans le fichier
# --------------------------------------------------------------------------

def bureaux_2022(raw: Path) -> pl.DataFrame:
    df = lire_csv(raw / "mi2022_t1_burvot.txt", latin1=True, truncate_ragged_lines=True)
    df = df.select(
        pl.col("Code du département").alias("dep_mi"),
        pl.col("Code de la circonscription").alias("num"),
        pl.col("Code de la commune").alias("com3"),
        pl.col("Libellé de la commune").alias("libelle"),
        pl.col("Code du b.vote").alias("bv"),
        pl.col("Inscrits").cast(pl.Int64).alias("inscrits"),
    )
    rows = []
    for r in df.iter_rows(named=True):
        dep = dep_an(r["dep_mi"])
        if r["dep_mi"] == "ZZ":
            insee = None
        elif r["dep_mi"] in PREFIXE_OM:
            insee = PREFIXE_OM[r["dep_mi"]] + r["com3"]
        else:
            insee = r["dep_mi"] + r["com3"]
        rows.append((dep, insee, insee, r["libelle"], r["bv"], r["inscrits"],
                     circo_id(dep, r["num"])))
    return pl.DataFrame(rows, schema=["dep", "insee", "insee_mi", "libelle", "bv", "inscrits",
                                      "circo"], orient="row")


def avec_arrondissements(df: pl.DataFrame) -> pl.DataFrame:
    """Remplace Paris / Marseille / Lyon par leurs arrondissements (codes Insee 751xx…)."""
    return df.with_columns(
        insee=pl.struct("insee", "bv").map_elements(
            lambda s: arrondissement(s["insee"], s["bv"]) if s["insee"] in PLM else s["insee"],
            return_dtype=pl.String,
        )
    )


def commune_circo(df: pl.DataFrame) -> pl.DataFrame:
    return (
        df.filter(pl.col("insee").is_not_null())
        .group_by("insee", "circo")
        .agg(pl.col("libelle").first(), pl.col("inscrits").sum(), pl.len().alias("nb_bv"))
        .with_columns(part=pl.col("inscrits") / pl.col("inscrits").sum().over("insee"))
        .sort("insee", "circo")
    )


# --------------------------------------------------------------------------
# COG : codes communes 2024 → codes en vigueur (fusions 2024-2026)
# --------------------------------------------------------------------------

def cog(raw: Path):
    com = pl.read_csv(raw / "cog2026_communes.csv", infer_schema=False)
    comer = pl.read_csv(raw / "cog2026_comer.csv", infer_schema=False)
    mvt = pl.read_csv(raw / "cog2026_mvt.csv", infer_schema=False)
    actuelles = set(com.filter(pl.col("TYPECOM").is_in(["COM", "ARM"]))["COM"]) | set(
        comer.filter(pl.col("NATURE_ZONAGE") == "COM")["COM_COMER"])
    # Dernier événement connu pour chaque ancienne commune disparue.
    m = (mvt.filter((pl.col("TYPECOM_AV") == "COM") & (pl.col("TYPECOM_AP") == "COM")
                    & (pl.col("COM_AV") != pl.col("COM_AP")))
         .sort("DATE_EFF").group_by("COM_AV").agg(pl.col("COM_AP").last(),
                                                   pl.col("DATE_EFF").last(),
                                                   pl.col("MOD").last()))
    suite = {r["COM_AV"]: (r["COM_AP"], r["DATE_EFF"]) for r in m.iter_rows(named=True)}

    def actuel(code: str) -> tuple[str, str | None]:
        vu, date = set(), None
        while code not in actuelles and code in suite and code not in vu:
            vu.add(code)
            code, date = suite[code]
        return code, date

    # Communes nouvelles : communes ayant au moins une commune déléguée en 2026,
    # ou issues d'une fusion depuis 2010.
    fusions = mvt.filter(pl.col("MOD").is_in(["31", "32", "33", "34"])
                         & (pl.col("DATE_EFF") >= "2010-01-01")
                         & (pl.col("COM_AV") != pl.col("COM_AP")))
    nouvelles = set(fusions["COM_AP"]) | set(
        com.filter(pl.col("TYPECOM") == "COMD")["COMPARENT"])
    deleguees = com.filter(pl.col("TYPECOM") == "COMD").select(
        pl.col("COM").alias("comd"), pl.col("LIBELLE").alias("lib_comd"),
        pl.col("COMPARENT").alias("insee"))
    vraies = com.filter(pl.col("TYPECOM").is_in(["COM", "ARM"]))  # pas les COMD homonymes
    libelles = dict(zip(vraies["COM"], vraies["LIBELLE"])) | dict(
        zip(comer["COM_COMER"], comer["LIBELLE"]))
    # Rétablissements (MOD 21) : commune recréée → commune dont elle sort.
    ret = mvt.filter((pl.col("MOD") == "21") & (pl.col("TYPECOM_AV") == "COM")
                     & (pl.col("TYPECOM_AP") == "COM") & (pl.col("COM_AV") != pl.col("COM_AP")))
    parents = dict(zip(ret["COM_AP"], ret["COM_AV"]))
    return actuelles, actuel, nouvelles, deleguees, libelles, parents


# --------------------------------------------------------------------------
# AMO : circonscription → député
# --------------------------------------------------------------------------

def deputes(amo: Path) -> pl.DataFrame:
    rows = []
    for f in glob.glob(str(amo / "*.json")):
        a = json.loads(Path(f).read_text())["acteur"]
        ms = a["mandats"]["mandat"]
        ms = ms if isinstance(ms, list) else [ms]
        ident = a["etatCivil"]["ident"]
        for m in ms:
            if m.get("typeOrgane") != "ASSEMBLEE" or m.get("legislature") not in ("16", "17"):
                continue
            lieu = (m.get("election") or {}).get("lieu") or {}
            rows.append({
                "uid": a["uid"]["#text"] if isinstance(a["uid"], dict) else a["uid"],
                "nom": f"{ident['prenom']} {ident['nom']}",
                "legislature": m["legislature"],
                "dep": lieu.get("numDepartement"), "libdep": lieu.get("departement"),
                "num": lieu.get("numCirco"),
                "debut": m.get("dateDebut"), "fin": m.get("dateFin"),
                "cause": (m.get("election") or {}).get("causeMandat"),
                "ref": (m.get("election") or {}).get("refCirconscription"),
            })
    df = pl.DataFrame(rows)
    return df.with_columns(
        circo=pl.concat_str([pl.col("dep"), pl.col("num").str.zfill(2)], separator="-"))


# --------------------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw", type=Path, default=RAW)
    ap.add_argument("--amo", type=Path, default=AMO)
    ap.add_argument("--out", type=Path, default=RAW / "derive")
    args = ap.parse_args()
    raw, out = args.raw, args.out
    out.mkdir(parents=True, exist_ok=True)
    pl.Config.set_tbl_rows(60)
    pl.Config.set_tbl_cols(12)
    pl.Config.set_fmt_str_lengths(60)

    # ---------------- Commune × circo, 2024 et 2022 ----------------
    titre("1. Commune × circonscription (ministère de l'Intérieur)")
    b24 = bureaux_2024(raw)
    b24 = b24.with_columns(
        insee=pl.col("com_mi").map_elements(code_insee_2024, return_dtype=pl.String))
    fe24 = b24.filter(pl.col("dep") == "099")
    b24 = avec_arrondissements(b24)
    cc24 = commune_circo(b24)
    b22 = avec_arrondissements(bureaux_2022(raw))
    cc22 = commune_circo(b22)
    for nom, cc in (("2024", cc24), ("2022", cc22)):
        print(f"{nom} : {cc['insee'].n_unique()} communes/arrondissements, "
              f"{cc.height} couples commune×circo, {cc['circo'].n_unique()} circos "
              f"(hors étranger)")
    anomalies = cc24.filter(pl.col("insee").str.ends_with("?"))
    print("Bureaux PLM hors numérotation par arrondissement (2024) :")
    print(anomalies)

    s24 = set(zip(cc24["insee"], cc24["circo"]))
    s22 = set(zip(cc22["insee"], cc22["circo"]))
    print(f"Couples présents en 2022 seulement : {len(s22 - s24)} ; en 2024 seulement : "
          f"{len(s24 - s22)}")
    print("  2022 seul :", sorted(s22 - s24))
    print("  2024 seul :", sorted(s24 - s22))

    # REU : contrôle du n° d'arrondissement déduit du n° de bureau (Paris/Lyon/Marseille).
    reu_bv = raw / "insee_reu_bv.parquet"
    if reu_bv.exists():
        r = pl.read_parquet(reu_bv).filter(
            pl.col("id_brut_miom").str.slice(0, 5).is_in(list(PLM)))
        r = r.with_columns(
            arr_bv=pl.col("id_brut_miom").map_elements(
                lambda s: arrondissement(s[:5], s[6:]), return_dtype=pl.String))
        ko = r.filter(pl.col("arr_bv") != pl.col("code_commune"))
        print(f"REU : {r.height} bureaux PLM, {ko.height} où l'arrondissement Insee "
              f"diffère du n° de bureau")
        if ko.height:
            print(ko.select("code_commune", "id_brut_miom", "arr_bv").head(10))

    # ---------------- Codes 2024 → COG 2026 ----------------
    titre("2. Communes nouvelles et fusions (COG 2026)")
    actuelles, actuel, nouvelles, deleguees, libelles, parents = cog(raw)
    cc24 = (cc24.with_columns(insee=pl.col("insee").map_elements(
        lambda c: CODES_SPECIAUX.get(c, [c]), return_dtype=pl.List(pl.String))).explode("insee"))
    m = cc24.with_columns(
        r=pl.col("insee").map_elements(lambda c: list(actuel(c)) if not c.endswith("?") else [c, None],
                                       return_dtype=pl.List(pl.String)))
    m = m.with_columns(insee_2026=pl.col("r").list.get(0), date_fusion=pl.col("r").list.get(1)).drop("r")
    changes = m.filter(pl.col("insee_2026") != pl.col("insee"))
    print(f"Codes commune 2024 remplacés depuis (fusions 2024-2026) : "
          f"{changes['insee'].n_unique()} → {changes['insee_2026'].n_unique()} communes")
    inconnus = m.filter(~pl.col("insee_2026").is_in(list(actuelles))
                        & ~pl.col("insee_2026").str.ends_with("?"))
    print(f"Codes 2024 absents du COG 2026 après suivi des fusions : {inconnus['insee'].n_unique()}")
    if inconnus.height:
        print(inconnus.select("insee", "libelle", "circo").head(20))

    cc = (m.filter(~pl.col("insee_2026").str.ends_with("?"))
          .group_by("insee_2026", "circo").agg(pl.col("inscrits").sum(), pl.col("nb_bv").sum())
          .rename({"insee_2026": "insee"})
          .with_columns(part=pl.col("inscrits") / pl.col("inscrits").sum().over("insee"),
                        libelle=pl.col("insee").replace_strict(libelles, default=None))
          .with_columns(source=pl.lit("MI 2024"))
          .sort("insee", "circo"))
    # Repli 1 : commune rétablie depuis 2024 → circo(s) de la commune dont elle sort.
    connues = set(cc["insee"])
    ajouts = []
    for code in sorted(actuelles - connues):
        if code in parents and parents[code] in connues:
            for c in cc.filter(pl.col("insee") == parents[code])["circo"]:
                ajouts.append((code, c, "rétablissement (COG)"))
    # Repli 2 : table MI 2017 (communes sans électeurs, etc.).
    t17 = table_2017(raw, actuel)
    if t17 is not None:
        connues |= {a[0] for a in ajouts}
        for code, c in sorted(t17):
            if code in actuelles and code not in connues and code not in PLM:
                ajouts.append((code, c, "MI 2017"))
    if ajouts:
        add = pl.DataFrame(ajouts, schema=["insee", "circo", "source"], orient="row")
        print(f"Communes ajoutées par repli : {add['insee'].n_unique()}")
        print(add.with_columns(libelle=pl.col("insee").replace_strict(libelles, default=None)))
        cc = pl.concat([cc, add], how="diagonal")
    cc.write_csv(out / "commune_circo_2026.csv")
    if t17 is not None:
        compare_2017(t17, cc)

    # Fusions postérieures à 2024 qui réunissent des communes de circos différentes.
    post = (changes.group_by("insee_2026").agg(pl.col("circo").unique().sort(),
                                               pl.col("insee").unique().sort(),
                                               pl.col("date_fusion").max())
            .filter(pl.col("circo").list.len() > 1))
    print(f"Fusions postérieures au scrutin de 2024 à cheval sur plusieurs circos : {post.height}")
    if post.height:
        print(post.with_columns(lib=pl.col("insee_2026").replace_strict(libelles, default=None)))

    # ---------------- Communes à cheval ----------------
    titre("3. Communes (et arrondissements PLM) à cheval sur plusieurs circonscriptions")
    nb = (cc.group_by("insee").agg(pl.col("libelle").first(), pl.len().alias("nb_circos"),
                                   pl.col("circo").sort().str.join(" "),
                                   pl.col("inscrits").sum(),
                                   pl.col("part").min().alias("part_min"))
          .sort("nb_circos", "inscrits", descending=True))
    multi = nb.filter(pl.col("nb_circos") > 1)
    print(f"Communes/arrondissements en vigueur couverts : {nb.height} ; "
          f"à cheval : {multi.height}")
    print(multi.group_by("nb_circos").len().sort("nb_circos"))
    multi = multi.with_columns(
        type=pl.when(pl.col("insee").str.slice(0, 3).is_in(["751", "132"])
                     | pl.col("insee").str.starts_with("6938")).then(pl.lit("arr. PLM"))
        .when(pl.col("insee").is_in(list(nouvelles))).then(pl.lit("commune nouvelle"))
        .otherwise(pl.lit("commune")))
    print(multi.group_by("type").len())
    print(multi.select("insee", "libelle", "type", "nb_circos", "circo", "inscrits", "part_min"))
    multi.write_csv(out / "communes_a_cheval.csv")
    tot = cc["inscrits"].sum()
    print(f"Inscrits 2024 vivant dans une commune/arrondissement à cheval : "
          f"{multi['inscrits'].sum():,} / {tot:,} ({multi['inscrits'].sum() / tot:.1%})")

    # Communes nouvelles à cheval : leurs communes déléguées, si le COG en garde.
    cn = multi.filter(pl.col("type") == "commune nouvelle")
    print(f"\nCommunes nouvelles à cheval : {cn.height} ; communes déléguées connues au COG :")
    print(cn.join(deleguees, on="insee", how="left").group_by("insee", "libelle", "circo")
          .agg(pl.col("lib_comd").drop_nulls().str.join(", ").alias("deleguees"))
          .sort("insee"))

    # Communes nouvelles à cheval : la coupure vient-elle de la fusion ?
    t12 = table_2012(raw)
    if t12 is not None:
        n12 = collections.Counter(c for c, _ in t12)
        avant = collections.defaultdict(set)
        for c in n12:
            avant[actuel(c)[0]].add(c)
        lignes = []
        for code, lib in zip(cn["insee"], cn["libelle"]):
            av = avant.get(code, set())
            deja = sorted(a for a in av if n12[a] > 1)
            lignes.append((code, lib, len(av), "commune d'origine déjà coupée en 2012" if deja
                           else "coupure due à la fusion"))
        v = pl.DataFrame(lignes, schema=["insee", "libelle", "communes_2012", "origine"],
                         orient="row")
        print("Origine de la coupure (table MI 2012) :")
        print(v.group_by("origine").agg(pl.len(), pl.col("libelle").head(12).str.join(", ")))

    # ---------------- Codes postaux ----------------
    titre("4. Codes postaux (La Poste) → circonscriptions")
    lp = lire_csv(raw / "laposte_hexasmal.csv", latin1=True)
    lp = lp.rename({lp.columns[0]: "insee", "Code_postal": "cp",
                    "Nom_de_la_commune": "nom", "Ligne_5": "ligne5"})
    print(f"La Poste : {lp.height} lignes, {lp['cp'].n_unique()} codes postaux, "
          f"{lp['insee'].n_unique()} codes commune")
    cp_com = lp.select("cp", "insee").unique()
    j = cp_com.join(cc.select("insee", "circo", "inscrits", "part"), on="insee", how="left")
    orphelins = j.filter(pl.col("circo").is_null()).join(
        lp.select("insee", "nom").unique("insee"), on="insee")
    print(f"Couples (cp, commune) sans circonscription : {orphelins.height}")
    print(orphelins.sort("insee"))
    j = j.filter(pl.col("circo").is_not_null())
    com_nb = cc.group_by("insee").agg(pl.len().alias("nb_circos_com"))
    j = j.join(com_nb, on="insee")
    par_cp = j.group_by("cp").agg(
        nb_communes=pl.col("insee").n_unique(),
        nb_circos=pl.col("circo").n_unique(),
        commune_coupee=(pl.col("nb_circos_com") > 1).any(),
        circos=pl.col("circo").unique().sort().str.join(" "),
    ).with_columns(
        cas=pl.when(pl.col("nb_circos") == 1).then(pl.lit("a. 1 circo"))
        .when(~pl.col("commune_coupee")).then(pl.lit("b. choisir la commune suffit"))
        .otherwise(pl.lit("c. commune coupée : adresse nécessaire")))
    print(f"Codes postaux rattachés : {par_cp.height}")
    print(par_cp.group_by("nb_circos").len().sort("nb_circos"))
    print(par_cp.group_by("cas").len().sort("cas"))
    print("Codes postaux à 1 seule commune :",
          par_cp.filter(pl.col("nb_communes") == 1).height,
          "; à plusieurs communes :", par_cp.filter(pl.col("nb_communes") > 1).height,
          "; max communes par CP :", par_cp["nb_communes"].max())
    print("CP aux circos les plus nombreuses :")
    print(par_cp.sort("nb_circos", descending=True).head(15))
    par_cp.sort("cp").write_csv(out / "cp_resume.csv")
    j.sort("cp", "insee", "circo").write_csv(out / "cp_commune_circo.csv")

    # ---------------- Adresses REU : (cp, commune) des communes coupées ----------------
    adr = raw / "insee_reu_adresses.parquet"
    if adr.exists():
        titre("5. Adresses du REU (Insee 2023) : le code postal suffit-il dans une commune coupée ?")
        reu_adresses(pl.scan_parquet(adr), b22, cc, lp, par_cp, actuel, out)

    # ---------------- Outre-mer et étranger ----------------
    titre("6. Outre-mer et Français de l'étranger")
    om = cc.filter(pl.col("insee").str.slice(0, 2).is_in(["97", "98"]))
    print(om.with_columns(dep=pl.col("circo").str.slice(0, 3))
          .group_by("dep").agg(communes=pl.col("insee").n_unique(),
                               circos=pl.col("circo").n_unique()).sort("dep"))
    print("CP La Poste outre-mer par préfixe :")
    print(lp.with_columns(p=pl.col("insee").str.slice(0, 3))
          .filter(pl.col("p") >= "970").group_by("p")
          .agg(pl.col("cp").n_unique().alias("cp"), pl.col("insee").n_unique().alias("communes"),
               pl.col("nom").unique().sort().head(3).str.join(", ").alias("exemples"))
          .sort("p"))
    print(f"Étranger (ZZ) 2024 : {fe24['circo'].n_unique()} circos, "
          f"{fe24['com_mi'].n_unique()} « communes » (postes consulaires), "
          f"{fe24.height} bureaux")
    print(fe24.group_by("circo").agg(pl.col("libelle").n_unique().alias("postes"),
                                     pl.col("libelle").unique().sort().head(5).str.join(", "))
          .sort("circo"))

    # ---------------- Jointure AN ----------------
    titre("7. Circonscription → député (AMO, XVIe et XVIIe)")
    d = deputes(args.amo)
    circos24 = set(b24["circo"].unique()) | set(fe24["circo"].unique())
    circos22 = set(b22["circo"].unique())
    print(f"Circos MI : 2024 = {len(circos24)}, 2022 = {len(circos22)}, "
          f"identiques : {circos24 == circos22}")
    for leg, ref in (("17", circos24), ("16", circos22)):
        dl = d.filter(pl.col("legislature") == leg)
        an = set(dl["circo"].unique())
        print(f"\nLégislature {leg} : {dl.height} mandats, {dl['uid'].n_unique()} députés, "
              f"{len(an)} circos")
        print("  circos MI sans mandat AN :", sorted(ref - an))
        print("  circos AN inconnues du MI :", sorted(an - ref))
        print("  codes département AN hors 2 chiffres :",
              sorted({x for x in dl["dep"].unique() if not (x.isdigit() and len(x) == 2)}))
        par = dl.group_by("circo").len()
        print("  mandats par circo :", dict(sorted(collections.Counter(par["len"]).items())))
        print("  causes :", dict(collections.Counter(dl["cause"])))
        if leg == "17":
            en_cours = dl.filter(pl.col("fin").is_null())
            print(f"  mandats en cours : {en_cours.height} sur {en_cours['circo'].n_unique()} circos")
            vac = sorted(ref - set(en_cours["circo"]))
            print("  circos sans député en cours (vacance ?) :", vac)
            for c in vac:
                print("   ", dl.filter(pl.col("circo") == c).sort("debut")
                      .select("nom", "debut", "fin", "cause").rows())
            print("  circos à plusieurs mandats en cours :",
                  en_cours.group_by("circo").len().filter(pl.col("len") > 1)["circo"].to_list())
            refs = dl.group_by("circo").agg(pl.col("ref").n_unique()).filter(pl.col("ref") > 1)
            print("  circos dont refCirconscription varie :", refs.height)


def reu_adresses(a: pl.LazyFrame, b22: pl.DataFrame, cc: pl.DataFrame, lp: pl.DataFrame,
                 par_cp: pl.DataFrame, actuel, out: Path) -> None:
    """Adresse REU → bureau REU → bureau MI 2022 → circo.

    1. Pour chaque couple (code postal, commune coupée), le code postal suffit-il ?
    2. Part des électeurs dans chaque cas d'UX. Le REU ne donne pas d'effectif par
       adresse : on répartit les inscrits 2024 de la commune au prorata de `nb_adresses`.
    """
    reu_bv = pl.read_parquet(RAW / "insee_reu_bv.parquet").select(
        pl.col("id_brut_reu").alias("id_brut_bv_reu"), "id_brut_miom")
    bv22 = b22.filter(pl.col("insee_mi").is_not_null()).select(
        id_brut_miom=pl.concat_str([pl.col("insee_mi"), pl.col("bv")], separator="_"),
        circo_bv=pl.col("circo"))
    adr = (a.select("code_commune_ref", "id_brut_bv_reu", "nb_adresses",
                    cp=pl.col("geo_adresse").str.extract(r"(\d{5})\D*$", 1))
           .group_by("code_commune_ref", "id_brut_bv_reu", "cp").agg(pl.col("nb_adresses").sum())
           .collect())
    codes = adr["code_commune_ref"].unique().to_list()
    vers = {c: actuel(c)[0] for c in codes}
    adr = adr.with_columns(insee=pl.col("code_commune_ref").replace_strict(vers))
    adr = adr.join(reu_bv, on="id_brut_bv_reu", how="left").join(bv22, on="id_brut_miom",
                                                                   how="left")
    nbc = cc.group_by("insee").agg(nb_circos_com=pl.len(), circo_unique=pl.col("circo").first(),
                                   inscrits_com=pl.col("inscrits").sum())
    adr = adr.join(nbc, on="insee", how="left").with_columns(
        circo=pl.when(pl.col("nb_circos_com") == 1).then(pl.col("circo_unique"))
        .otherwise(pl.col("circo_bv")),
        electeurs=pl.col("inscrits_com") * pl.col("nb_adresses")
        / pl.col("nb_adresses").sum().over("insee"))
    tot = adr["electeurs"].sum()
    coup = adr.filter(pl.col("nb_circos_com") > 1)
    print(f"REU : {adr.height:,} lignes (commune, bureau, CP) ; communes sans circo : "
          f"{adr.filter(pl.col('nb_circos_com').is_null())['insee'].n_unique()} ; "
          f"adresses des communes coupées rattachées à un bureau MI 2022 : "
          f"{coup.filter(pl.col('circo').is_not_null())['nb_adresses'].sum() / coup['nb_adresses'].sum():.1%}")
    lpc = lp.select("cp", "insee").unique()
    adr_lp = adr.join(lpc, on=["cp", "insee"], how="semi")
    print(f"Part des électeurs (estimés) dont le CP géocodé est un couple (CP, commune) "
          f"La Poste : {adr_lp['electeurs'].sum() / tot:.1%}")

    g = (adr_lp.filter(pl.col("circo").is_not_null())
         .group_by("cp", "insee", "circo").agg(pl.col("nb_adresses").sum())
         .with_columns(part=pl.col("nb_adresses") / pl.col("nb_adresses").sum().over("cp", "insee")))
    # Rattachements marginaux (< 2 % des adresses du couple) ignorés : bruit de géocodage.
    pc = (g.filter(pl.col("part") >= 0.02).group_by("cp", "insee")
          .agg(nb_circos=pl.len(), circos=pl.col("circo").sort().str.join(" "),
               part_max=pl.col("part").max(), adresses=pl.col("nb_adresses").sum()))
    pcc = pc.join(cc.group_by("insee").agg(pl.len().alias("n")).filter(pl.col("n") > 1),
                  on="insee", how="semi")
    print(f"\nCouples (CP, commune coupée) : {pcc.height} ; nombre de circos possibles :")
    print(pcc.group_by("nb_circos").agg(pl.len().alias("couples"),
                                         pl.col("adresses").sum()).sort("nb_circos"))
    print("Exemples : Toulouse, Nice, Montpellier, Bordeaux, Paris 15e, Marseille 13e")
    print(pcc.filter(pl.col("insee").is_in(
        ["31555", "06088", "34172", "33063", "75115", "13213"])).sort("insee", "cp"))
    pcc.sort("insee", "cp").write_csv(out / "cp_commune_coupee_reu.csv")

    # Classement final des CP et part des électeurs.
    ok = pc.group_by("cp").agg(tous_uniques=(pl.col("nb_circos") == 1).all())
    cl = par_cp.join(ok, on="cp", how="left").with_columns(
        cas2=pl.when(pl.col("cas").str.starts_with("c") & pl.col("tous_uniques").fill_null(False))
        .then(pl.lit("c1. CP + commune suffit (le CP ne touche qu'une circo de la commune)"))
        .when(pl.col("cas").str.starts_with("c")).then(pl.lit("c2. adresse nécessaire"))
        .otherwise(pl.col("cas")))
    el = adr_lp.group_by("cp").agg(pl.col("electeurs").sum())
    res = (cl.join(el, on="cp", how="left").group_by("cas2")
           .agg(cp=pl.len(), electeurs=pl.col("electeurs").sum())
           .with_columns(part_electeurs=pl.col("electeurs") / pl.col("electeurs").sum())
           .sort("cas2"))
    print("\nClassement des codes postaux (électeurs estimés via REU) :")
    print(res)
    cl.sort("cp").write_csv(out / "cp_resume.csv")
    # À l'échelle de l'électeur qui a choisi sa commune : reste-t-il plusieurs circos ?
    elc = adr_lp.group_by("cp", "insee").agg(pl.col("electeurs").sum()).join(
        pc.select("cp", "insee", "nb_circos"), on=["cp", "insee"], how="left")
    amb = elc.filter(pl.col("nb_circos") > 1)["electeurs"].sum()
    print(f"Électeurs (estimés) dont le couple (CP, commune) reste ambigu : {amb:,.0f} "
          f"({amb / elc['electeurs'].sum():.1%})")
    amb_c = elc.join(cc.group_by("insee").agg(pl.len().alias("n")).filter(pl.col("n") > 1),
                     on="insee", how="semi")["electeurs"].sum()
    print(f"Électeurs (estimés) dans une commune/arrondissement coupé : {amb_c:,.0f} "
          f"({amb_c / elc['electeurs'].sum():.1%})")


def table_2012(raw: Path) -> set[tuple[str, str]] | None:
    """Table MI 2012 (communes 2012 × circos 2012), hors outre-mer et étranger."""
    try:
        import xlrd
    except ImportError:
        print("(table MI 2012 ignorée : lancer avec `uv run --with xlrd`)")
        return None
    sh = xlrd.open_workbook(raw / "mi2012_correspondance.xls").sheet_by_index(0)
    rows = set()
    for i in range(2, sh.nrows):
        r = sh.row_values(i)
        dep, com, circ = r[1], r[5], r[8]
        dep = f"{int(dep):02d}" if isinstance(dep, float) else str(dep).strip()
        if not com or not circ or len(dep) != 2:
            continue
        com = str(int(com)) if isinstance(com, float) else str(com).strip()
        circ = str(int(circ)) if isinstance(circ, float) else str(circ).strip()
        if not (com.isdigit() and circ.isdigit()):
            continue
        rows.add((dep + com.zfill(3), circo_id(dep, circ)))
    return rows




def table_2017(raw: Path, actuel) -> set[tuple[str, str]] | None:
    """Table MI 2017 (communes × cantons × circos), codes ramenés au COG 2026."""
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("\n(table MI 2017 ignorée : lancer avec `uv run --with openpyxl`)")
        return None
    t = pl.read_excel(raw / "mi2017_correspondance.xlsx", engine="openpyxl",
                      infer_schema_length=0)
    t = t.select(pl.all().cast(pl.String))
    rows = set()
    for r in t.iter_rows(named=True):
        dep_mi, com, num = r["CODE DPT"], r["CODE COMMUNE"].zfill(3), r["CODE CIRC LEGISLATIVE"]
        if dep_mi == "ZZ":
            continue
        if dep_mi == "ZM" and com.startswith("5"):  # anciens codes de Mayotte (985xx)
            com = "6" + com[1:]
        if dep_mi == "ZP" and com.startswith("0"):  # Polynésie codée sans le « 7 »
            com = "7" + com[1:]
        dep = dep_an(dep_mi)
        insee = (PREFIXE_OM[dep_mi] + com) if dep_mi in PREFIXE_OM else dep_mi.zfill(2) + com
        for code in (["98611", "98612", "98613"] if dep_mi == "ZW" else [insee]):
            rows.add((actuel(code)[0], circo_id(dep, num)))
    return rows


def compare_2017(rows: set[tuple[str, str]], cc: pl.DataFrame) -> None:
    titre("2 bis. Table de correspondance MI 2017 vs résultats 2024 (codes COG 2026)")
    s26 = set(zip(cc.filter(pl.col("source") == "MI 2024")["insee"],
                  cc.filter(pl.col("source") == "MI 2024")["circo"]))
    # PLM : la table 2017 est à la commune ; on la compare hors PLM.
    hors = {x for x in s26 if not x[0].startswith(("751", "132", "6938"))}
    rows = {x for x in rows if x[0] not in PLM}
    print(f"2017 → COG 2026 : {len(rows)} couples ; 2024 → COG 2026 (hors PLM) : {len(hors)}")
    seul17, seul24 = sorted(rows - hors), sorted(hors - rows)
    print(f"Couples dans 2017 seulement : {len(seul17)} ; dans 2024 seulement : {len(seul24)}")
    print("  2017 seul :", seul17)
    print("  2024 seul :", seul24)


if __name__ == "__main__":
    main()
